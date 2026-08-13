#!/usr/bin/env python3
"""
Extract Sheet1 from the source workbook into a dictionary-encoded JSON payload
for embedding in rpa-monitor.html.

Sheet1 is the ONLY sheet read. Every other sheet in the workbook is a derived
or specification sheet and is never a runtime input.

Why dictionary encoding: the log carries ~2,000 distinct free-text failure
reasons against only ~70 distinct vendor names, and reason strings dominate the
payload. Interning both into shared arrays and storing integer indices per row
roughly halves the embedded size (2.89 MB -> 1.29 MB) with no information loss.

The row validation here mirrors, exactly, what the dashboard does at runtime for
a user-supplied file. It is duplicated deliberately: the browser must be able to
reach the same verdict on a replacement dataset without this script.

    python3 tools/extract_sheet1.py data/source_workbook.xlsx data/sheet1_payload.json
"""
from __future__ import annotations

import datetime as dt
import json
import sys

import openpyxl

SHEET = "Sheet1"
FIELDS = ("Status", "Vendor", "Reason", "PO", "Date")

# Row outcomes. OK rows reach the engines; everything else is quarantined and
# reported in the dashboard's data-quality panel, never silently dropped.
OK = "OK"
BLANK_ROW = "EXCLUDED_BLANK_ROW"
NO_VENDOR = "EXCLUDED_NO_VENDOR"
BAD_DATE = "EXCLUDED_BAD_DATE"


def _text(value) -> str:
    """Excel cell -> trimmed string. None and numeric cells both survive."""
    if value is None:
        return ""
    return str(value).strip()


def _iso(value):
    """Excel cell -> naive ISO 8601 string, or None when it is not a datetime.

    Timestamps in this log carry no timezone and none is implied, so they are
    emitted without an offset and parsed back as local time in the browser.
    """
    if isinstance(value, dt.datetime):
        return value.isoformat()
    if isinstance(value, dt.date):
        return dt.datetime(value.year, value.month, value.day).isoformat()
    return None


def classify(status: str, vendor: str, reason: str, po: str, date_iso, raw) -> str:
    """Decide a row's fate. Order matters: fully-blank is reported separately
    from a real transaction that merely lost its vendor, because the two mean
    very different things to whoever maintains the source system."""
    if not any(_text(v) for v in raw):
        return BLANK_ROW
    if date_iso is None:
        return BAD_DATE
    if not vendor:
        return NO_VENDOR
    return OK


def extract(path: str) -> dict:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        raise SystemExit(f"workbook has no {SHEET!r} sheet (found: {wb.sheetnames})")

    rows = wb[SHEET].iter_rows(values_only=True)
    header = [_text(h) for h in next(rows)]
    try:
        idx = {f: header.index(f) for f in FIELDS}
    except ValueError as exc:
        raise SystemExit(f"Sheet1 is missing a required column: {exc}") from exc

    # The workbook's own "Total Transactions" counts any row with content in ANY
    # column, including the derived Category column. Sheet1's Category formula
    # spills past the data and leaves #VALUE! on rows that hold no transaction
    # at all, which inflates that figure. We reproduce it purely so the dashboard
    # can show the workbook's number beside the true one and explain the gap.
    workbook_total = 0
    workbook_failures = 0

    statuses: list[str] = []
    vendors: list[str] = []
    reasons: list[str] = []
    status_ix: dict[str, int] = {}
    vendor_ix: dict[str, int] = {}
    reason_ix: dict[str, int] = {}

    def intern(value: str, pool: list[str], index: dict[str, int]) -> int:
        if value not in index:
            index[value] = len(pool)
            pool.append(value)
        return index[value]

    packed: list[list] = []
    quarantine: list[dict] = []
    counts = {OK: 0, BLANK_ROW: 0, NO_VENDOR: 0, BAD_DATE: 0}
    non_empty = 0
    non_empty_failures = 0
    raw_total = 0
    # Non-Success/Failure status values are kept (is_failure is simply false),
    # matching the workbook's case-insensitive = "Failure" test, but they are
    # counted so an unexpected vocabulary in a new dataset stays visible.
    odd_status: dict[str, int] = {}

    for row_number, raw in enumerate(rows, start=2):
        raw_total += 1
        if any(_text(c) for c in raw):
            workbook_total += 1
            if _text(raw[idx["Status"]]).casefold() == "failure":
                workbook_failures += 1
        cells = [raw[idx[f]] if idx[f] < len(raw) else None for f in FIELDS]
        status, vendor, reason, po, date_cell = cells
        status_s, vendor_s = _text(status), _text(vendor)
        reason_s = "" if reason is None else str(reason)
        po_s = _text(po)
        date_iso = _iso(date_cell)

        verdict = classify(status_s, vendor_s, reason_s, po_s, date_iso, cells)
        counts[verdict] += 1

        if verdict != BLANK_ROW:
            non_empty += 1
            if status_s.casefold() == "failure":
                non_empty_failures += 1
            elif status_s.casefold() != "success":
                odd_status[status_s] = odd_status.get(status_s, 0) + 1

        if verdict != OK and verdict != BLANK_ROW:
            quarantine.append({
                "row": row_number,
                "verdict": verdict,
                "Status": status_s,
                "Vendor": vendor_s,
                "Reason": reason_s[:300],
                "PO": po_s,
                "Date": date_iso or _text(date_cell),
            })

        # EVERY row is packed, including the ones this script would reject.
        # The dashboard runs its own validation over the same rules, so the
        # reconciliation it shows is derived from the data in front of it rather
        # than inherited from here — and the validation path gets exercised
        # against real dirty rows instead of only synthetic ones.
        packed.append([
            intern(status_s, statuses, status_ix),
            intern(vendor_s, vendors, vendor_ix),
            intern(reason_s, reasons, reason_ix),
            po_s,
            date_iso if date_iso is not None else _text(date_cell),
        ])

    return {
        "generated": dt.datetime.now().isoformat(timespec="seconds"),
        "source": path,
        "statuses": statuses,
        "vendors": vendors,
        "reasons": reasons,
        "rows": packed,
        "quality": {
            "rawRows": raw_total,
            "workbookTotal": workbook_total,
            "workbookFailures": workbook_failures,
            "nonEmptyRows": non_empty,
            "nonEmptyFailures": non_empty_failures,
            "analysable": counts[OK],
            "excludedBlankRow": counts[BLANK_ROW],
            "excludedNoVendor": counts[NO_VENDOR],
            "excludedBadDate": counts[BAD_DATE],
            "quarantine": quarantine,
            "oddStatus": dict(sorted(odd_status.items(), key=lambda kv: -kv[1])),
        },
    }


def main() -> None:
    src = sys.argv[1] if len(sys.argv) > 1 else "data/source_workbook.xlsx"
    dst = sys.argv[2] if len(sys.argv) > 2 else "data/sheet1_payload.json"
    payload = extract(src)
    with open(dst, "w", encoding="utf-8") as fh:
        json.dump(payload, fh, separators=(",", ":"), ensure_ascii=False)

    q = payload["quality"]
    size_mb = len(json.dumps(payload, separators=(",", ":"), ensure_ascii=False).encode()) / 1e6
    print(f"wrote {dst}  ({size_mb:.2f} MB)")
    print(f"  raw rows          {q['rawRows']:>7,}")
    print(f"  workbook total    {q['workbookTotal']:>7,}   <- workbook 'Total Transactions' "
          f"(rate {q['workbookFailures'] / q['workbookTotal']:.16f})")
    print(f"  real transactions {q['nonEmptyRows']:>7,}   "
          f"rate {q['nonEmptyFailures'] / q['nonEmptyRows']:.10f}")
    print(f"  analysable        {q['analysable']:>7,}")
    print(f"  excluded blank    {q['excludedBlankRow']:>7,}")
    print(f"  excluded no vendor{q['excludedNoVendor']:>7,}")
    print(f"  excluded bad date {q['excludedBadDate']:>7,}")
    print(f"  distinct vendors  {len(payload['vendors']):>7,}")
    print(f"  distinct reasons  {len(payload['reasons']):>7,}")


if __name__ == "__main__":
    main()
